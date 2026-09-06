"""Application-only document ingestion.

Text chunking mirrors ``src/ingest.py`` (page-wise cleanup, 500-word chunks,
100-word overlap). Tables are extracted separately and chunked with headers
retained. Research ingest.py is not used by the application path.
"""

from __future__ import annotations

import hashlib
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional


CHUNK_SIZE_WORDS = 500
CHUNK_OVERLAP_WORDS = 100
CHUNK_STRIDE_WORDS = CHUNK_SIZE_WORDS - CHUNK_OVERLAP_WORDS
TABLE_ROW_GROUP = 8


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str


@dataclass(frozen=True)
class ExtractedTable:
    page_number: int
    table_index: int
    title: str
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    section: str = ""


@dataclass(frozen=True)
class AppChunk:
    document_id: str
    document: str
    access_level: str
    chunk_id: int
    page_number: int
    text: str
    content_type: str = "text"
    table_index: Optional[int] = None
    table_title: str = ""
    section: str = ""

    @property
    def chroma_id(self) -> str:
        kind = "t" if self.content_type == "table" else "p"
        extra = self.table_index if self.table_index is not None else self.page_number
        return f"app_{self.document_id}_{kind}{extra}_c{self.chunk_id}"

    def metadata(self) -> dict[str, Any]:
        meta: dict[str, Any] = {
            "scope": "application",
            "document_id": self.document_id,
            "document": self.document,
            "access_level": self.access_level,
            "page": self.page_number,
            "chunk_id": self.chunk_id,
            "content_type": self.content_type,
        }
        if self.content_type == "table":
            meta["table_index"] = int(self.table_index or 0)
            meta["table_title"] = self.table_title or ""
            meta["section"] = self.section or ""
        return meta

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_id": self.document_id,
            "document": self.document,
            "access_level": self.access_level,
            "chunk_id": self.chunk_id,
            "page_number": self.page_number,
            "text": self.text,
            "content_type": self.content_type,
            "table_index": self.table_index,
            "table_title": self.table_title,
            "section": self.section,
        }


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def serialize_table(table: ExtractedTable) -> str:
    headers = [h for h in table.headers]
    rows = [list(row) for row in table.rows]
    title = table.title or "Untitled table"
    lines = [
        f"TABLE: {title}",
        f"PAGE: {table.page_number}",
        f"TABLE_INDEX: {table.table_index}",
    ]
    if table.section:
        lines.append(f"SECTION: {table.section}")
    if headers:
        lines.append("| " + " | ".join(headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in rows:
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        lines.append("| " + " | ".join(padded[: len(headers) or len(padded)]) + " |")
    statements = []
    for row in rows:
        if not headers:
            statements.append(" ".join(row).strip())
            continue
        pairs = []
        for header, value in zip(headers, row):
            if header and value:
                pairs.append(f"{header} is {value}")
        if pairs:
            statements.append(f"{title}. " + " and ".join(pairs) + ".")
    if statements:
        lines.append("")
        lines.extend(statements)
    return "\n".join(lines).strip()


def _rows_from_grid(grid: list[list[Any]]) -> ExtractedTable | None:
    cleaned = [[_cell(cell) for cell in row] for row in grid if any(_cell(cell) for cell in row)]
    if len(cleaned) < 2:
        return None
    width = max(len(row) for row in cleaned)
    if width < 2:
        return None
    normalized = [tuple((row + [""] * width)[:width]) for row in cleaned]
    headers = normalized[0]
    rows = tuple(normalized[1:])
    title = headers[0] if headers[0] and all(not h for h in headers[1:]) else " ".join(h for h in headers if h)[:80]
    return ExtractedTable(page_number=1, table_index=1, title=title or "Table", headers=headers, rows=rows)


def _extract_txt_tables(text: str) -> list[ExtractedTable]:
    tables: list[ExtractedTable] = []
    lines = text.splitlines()
    block: list[str] = []
    index = 1

    def flush():
        nonlocal index, block
        grid = []
        for line in block:
            if re.match(r"^\s*\|?\s*-{3,}", line):
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if any(cells):
                grid.append(cells)
        table = _rows_from_grid(grid)
        if table:
            tables.append(
                ExtractedTable(
                    page_number=1,
                    table_index=index,
                    title=table.title,
                    headers=table.headers,
                    rows=table.rows,
                )
            )
            index += 1
        block = []

    for line in lines:
        if "|" in line and line.count("|") >= 2:
            block.append(line)
        elif block:
            flush()
    if block:
        flush()
    return tables


def extract_pdf_structure(raw: bytes) -> tuple[list[ExtractedPage], list[ExtractedTable]]:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    document = pymupdf.open(stream=raw, filetype="pdf")
    pages: list[ExtractedPage] = []
    tables: list[ExtractedTable] = []
    table_index = 1
    try:
        for page_number, page in enumerate(document, start=1):
            text = clean_text(page.get_text("text"))
            if text:
                pages.append(ExtractedPage(page_number, text))
            try:
                finder = page.find_tables()
                found = list(finder.tables) if finder is not None else []
            except Exception:
                found = []
            for table in found:
                try:
                    grid = table.extract()
                except Exception:
                    continue
                parsed = _rows_from_grid(grid or [])
                if not parsed:
                    continue
                tables.append(
                    ExtractedTable(
                        page_number=page_number,
                        table_index=table_index,
                        title=parsed.title,
                        headers=parsed.headers,
                        rows=parsed.rows,
                    )
                )
                table_index += 1
    finally:
        document.close()
    return pages, tables


def extract_docx_structure(raw: bytes) -> tuple[list[ExtractedPage], list[ExtractedTable]]:
    try:
        from docx import Document
    except ImportError as error:
        raise ValueError("DOCX uploads require python-docx. Install requirements.txt and restart.") from error
    import io

    document = Document(io.BytesIO(raw))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    text = clean_text("\n".join(paragraphs))
    pages = [ExtractedPage(1, text)] if text else []
    tables: list[ExtractedTable] = []
    heading = ""
    for paragraph in document.paragraphs:
        style = (paragraph.style.name or "") if paragraph.style else ""
        if "Heading" in style and paragraph.text.strip():
            heading = paragraph.text.strip()
    for index, table in enumerate(document.tables, start=1):
        grid = [[_cell(cell.text) for cell in row.cells] for row in table.rows]
        parsed = _rows_from_grid(grid)
        if not parsed:
            continue
        tables.append(
            ExtractedTable(
                page_number=1,
                table_index=index,
                title=parsed.title or heading or f"Table {index}",
                headers=parsed.headers,
                rows=parsed.rows,
                section=heading,
            )
        )
    return pages, tables


def extract_pages(filename: str, raw: bytes) -> list[ExtractedPage]:
    pages, _tables = extract_document(filename, raw)
    return pages


def extract_document(filename: str, raw: bytes) -> tuple[list[ExtractedPage], list[ExtractedTable]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        decoded = raw.decode("utf-8", errors="replace")
        tables = _extract_txt_tables(decoded)
        text = clean_text(decoded)
        pages = [ExtractedPage(1, text)] if text else []
        return pages, tables
    if suffix == ".pdf":
        return extract_pdf_structure(raw)
    if suffix == ".docx":
        return extract_docx_structure(raw)
    raise ValueError("Only PDF, DOCX, and TXT uploads are supported.")


def chunk_pages(
    *,
    document_id: str,
    document_name: str,
    access_level: str,
    pages: list[ExtractedPage],
    start_chunk_id: int = 1,
) -> list[AppChunk]:
    chunks: list[AppChunk] = []
    next_chunk_id = start_chunk_id
    for page in pages:
        words = page.text.split()
        start = 0
        while start < len(words):
            chunk_words = words[start : start + CHUNK_SIZE_WORDS]
            if not chunk_words:
                break
            chunks.append(
                AppChunk(
                    document_id=document_id,
                    document=document_name,
                    access_level=access_level,
                    chunk_id=next_chunk_id,
                    page_number=page.page_number,
                    text=" ".join(chunk_words),
                    content_type="text",
                )
            )
            next_chunk_id += 1
            start += CHUNK_STRIDE_WORDS
    return chunks


def chunk_tables(
    *,
    document_id: str,
    document_name: str,
    access_level: str,
    tables: list[ExtractedTable],
    start_chunk_id: int = 1,
) -> list[AppChunk]:
    chunks: list[AppChunk] = []
    next_chunk_id = start_chunk_id
    for table in tables:
        rows = list(table.rows)
        groups = [rows] if len(rows) <= TABLE_ROW_GROUP else [rows[i : i + TABLE_ROW_GROUP] for i in range(0, len(rows), TABLE_ROW_GROUP)]
        for group in groups:
            subset = ExtractedTable(
                page_number=table.page_number,
                table_index=table.table_index,
                title=table.title,
                headers=table.headers,
                rows=tuple(group),
                section=table.section,
            )
            chunks.append(
                AppChunk(
                    document_id=document_id,
                    document=document_name,
                    access_level=access_level,
                    chunk_id=next_chunk_id,
                    page_number=table.page_number,
                    text=serialize_table(subset),
                    content_type="table",
                    table_index=table.table_index,
                    table_title=table.title,
                    section=table.section,
                )
            )
            next_chunk_id += 1
    return chunks


def build_chunks(*, document_id: str, document_name: str, access_level: str, pages: list[ExtractedPage], tables: list[ExtractedTable] | None = None) -> list[AppChunk]:
    text_chunks = chunk_pages(document_id=document_id, document_name=document_name, access_level=access_level, pages=pages)
    table_chunks = chunk_tables(
        document_id=document_id,
        document_name=document_name,
        access_level=access_level,
        tables=tables or [],
        start_chunk_id=len(text_chunks) + 1,
    )
    return text_chunks + table_chunks


def validate_chunks(chunks: list[AppChunk]) -> None:
    if not chunks:
        raise ValueError("Document extraction failed.")
    ids = [chunk.chroma_id for chunk in chunks]
    if len(ids) != len(set(ids)):
        raise ValueError("Generated chunk IDs are not unique.")
    for chunk in chunks:
        if not chunk.document_id or not chunk.document:
            raise ValueError("Generated chunk is missing document identity.")
        if chunk.page_number < 1:
            raise ValueError("Generated chunk has an invalid page number.")
        if not chunk.text.strip():
            raise ValueError("Generated chunk is empty.")
        if chunk.content_type == "table" and chunk.table_index is None:
            raise ValueError("Table chunk is missing table_index.")


def validate_embeddings(embeddings: Any, expected_count: int) -> list[list[float]]:
    if hasattr(embeddings, "tolist"):
        embeddings = embeddings.tolist()
    vectors = list(embeddings or [])
    if len(vectors) != expected_count:
        raise ValueError(f"Embedding generation failed: expected {expected_count} embeddings, got {len(vectors)}.")
    if not vectors:
        raise ValueError("Embedding generation failed.")
    dimension = len(vectors[0])
    if dimension <= 0:
        raise ValueError("Generated embeddings are empty.")
    for vector in vectors:
        if len(vector) != dimension:
            raise ValueError("Generated embeddings have inconsistent dimensions.")
        for value in vector:
            number = float(value)
            if not math.isfinite(number):
                raise ValueError("Generated embeddings contain invalid numeric values.")
    return vectors


def verify_chroma_records(collection: Any, ids: list[str]) -> int:
    if not hasattr(collection, "get"):
        return len(ids)
    result = collection.get(ids=ids)
    found = result.get("ids", []) if isinstance(result, dict) else []
    if len(found) != len(ids):
        raise ValueError(f"Vector indexing failed: expected {len(ids)} Chroma records, found {len(found)}.")
    return len(found)
